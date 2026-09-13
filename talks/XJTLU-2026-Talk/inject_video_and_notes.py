"""Inject the 36s student-games reel into slide 12 and add speaker notes to every slide."""

from __future__ import annotations
from pathlib import Path

import copy

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

PROJ = Path("/Users/xz33/WorkBuddy/来自生活的游戏设计/talks/XJTLU-2026-Talk")
PPTX = PROJ / "XJTLU-2026-Talk.pptx"
VIDEO = PROJ / "assets" / "student-games-reel.mp4"
POSTER = PROJ / "assets" / "30plus-games-contact.png"

# 1280x720 canvas. Place video centered, 960x540 (16:9), 160/90 margins.
LEFT = Emu(160 * 9525)
TOP = Emu(90 * 9525)
WIDTH = Emu(960 * 9525)
HEIGHT = Emu(540 * 9525)

NOTES = {
    1: (
        "Good morning. The title gives away the argument: from lived experience to playable systems, "
        "and the question of what AI does, and does not, change in that journey. "
        "I have ten minutes. I will spend the first two on the one-prompt trend and its empirical diagnosis, "
        "two on our setting, three on the method, and the rest on what two summer programs showed us — "
        "ending with one sentence I would like you to remember."
    ),
    2: (
        "Three chapters. The first frames the problem with recent industry and academic numbers. "
        "The second is the method: Game Design from Life in six steps, then the move from prompt to guided interview, "
        "and the development log as accountability. The third is what we saw, and where it points."
    ),
    3: (
        "Chapter one. I want to set the scene in three beats: a short industry trend, a short empirical counterpoint, "
        "and the specific classroom that I am in. Each of those will get its own slide."
    ),
    4: (
        "From late 2024 onward, generating a playable game from a one-sentence prompt has become a standard way "
        "to demonstrate model capability. Rosebud, Summer Engine, chat-driven IDEs, and the engine vendors themselves "
        "all lean in. GDC 2026 puts adoption over fifty percent. "
        "But the number on the right is a warning: zero of ten thousand four hundred single-pass generations "
        "compiled without a fix loop. Democratization of the demo is real — the question is what survives the demo."
    ),
    5: (
        "Four numbers from four independent 2026 papers. The mechanism fidelity F1 of direct natural-language to C# "
        "is statistically indistinguishable from doing nothing — about 0.12. The runtime pass rate collapses from "
        "eighty percent on small projects to under six percent on real ones. The best coding agent on an end-to-end "
        "benchmark reaches forty-one percent. The single largest proven gain is not a better prompt — it is putting "
        "a GUI-agent playtester in the loop, plus thirty-seven points on the rubric. "
        "Failure lives in structure and architecture, not in syntax. This is the gap our pedagogy has to address."
    ),
    6: (
        "So why teach game design at a research university in 2026? Our IMA students are not training to be "
        "game developers. They come from across the liberal-arts curriculum, with mixed coding experience, and they "
        "enroll because they have something they want to express — a frustration, a memory, a practice. "
        "Our mission is to teach them to use a system to express an experience. The success question is not "
        "is it fun, but did the translation succeed. Because coding is not the learning goal, AI assistance is "
        "legitimate. Because intent and emotional structure are the learning goal, those cannot be outsourced. "
        "The course lives exactly in that boundary, and that is the boundary I want to talk about."
    ),
    7: (
        "Chapter two. The method has three pieces. A six-step structure that takes a lived experience to a playable "
        "system. The decisive move from prompt to guided interview. And the development log as accountability. "
        "I will spend a slide on each."
    ),
    8: (
        "Six steps. Step one is a specific moment, not a topic. Step two is its emotional structure — trigger, tension, "
        "turn, peak, aftertaste. Step three datalizes the structure into named, bounded, tunable variables. "
        "Step three-point-five turns the variables into an inspectable system graph, with proposal and decision provenance. "
        "Step four is a minimum playable prototype — in our case, on the Tuanjie Engine, which is what NYU Shanghai students use. "
        "Steps five and six are where AI enters, and where the world is expanded around the core mechanic. "
        "The thing I want to underline is that AI enters at step five, not at step one."
    ),
    9: (
        "This is the slide I care about most. On the left, a one-prompt user compresses an entire emotional target "
        "into a few words, gets one shot, and has no trace. On the right, a guided interview unfolds the same target "
        "across many turns, with reflection pauses, and produces not a finished game but a system description — "
        "intent, mechanics, variables, feedback — that a human can read and that an AI can implement. "
        "Intent is not compressed in five words. It is unfolded by a question list until it is verifiable. "
        "This is the move from prompt to system description, and it is the heart of the pedagogy."
    ),
    10: (
        "Once AI is in the loop, the question becomes: who decided what? The development log records each interaction "
        "as a three-line trace: what the student asked for, a short summary of what the agent produced, "
        "and the student's decision — accepted, rejected, rewritten, and why. "
        "In The Darkroom, this trace runs seventy-two times. In Coaster Carnage, the student replaced the AI's "
        "obstacle spawner with a prebuilt track; the cart moves forward and the player only reacts. The student's "
        "reflection says explicitly that the new architecture felt more like being carried by the coaster than driving it. "
        "Without the log, that decision is invisible. With it, it is the artifact we study."
    ),
    11: (
        "Chapter three. I will play a thirty-six-second reel of twelve student works. Then I will show one case in depth, "
        "name three findings, and close with where I think the field should go next."
    ),
    12: (
        "This is the moment to stop talking. Thirty-six seconds, twelve works, drawn from this summer's high-school camp "
        "in Shanghai. Zero of these students had written a line of code before. Each picked a domain they actually knew — "
        "skating, climbing, curling, golf, archery, cello, crochet, first diagnosis, virtual kitchen, paper folding — "
        "and translated it into a playable system. Please watch, and remember what each game is trying to feel like."
    ),
    13: (
        "The Darkroom, by Boyan. The lived experience is the frustration of a long-exposure photograph that comes back "
        "wrong. The data are four variables — Exposure State, Light Trail, Trail Duration, Game Time. The build has "
        "fifty-one Unity scripts, and the log records seventy-two interactions. The system graph on the left is the "
        "artifact I want to point to: it shows intent, data, and player input in one place, in language a human can read. "
        "That is the system description our pedagogy produces. The student quote at the bottom right is the kind of "
        "decision a chat history alone would never let us see."
    ),
    14: (
        "Three findings. Feasible — zero-coding students can produce a playable demo in days. Asymmetric verification — "
        "AI can tell you the system runs, it cannot tell you the prototype still carries the feeling. That is the part "
        "where a playtest loop and a human reader are indispensable, and it echoes Väkevä and colleagues' CHI 2026 "
        "Best Paper argument that resonance with lived experience is the strongest anchor of meaningful gameplay. "
        "And learning happens in the experiment — at the cost of too many questions and not enough real-time feedback."
    ),
    15: (
        "Three trends. Lower floor to a playable demo. A growing case library of system descriptions and process logs. "
        "And a clear demand for specialized, one-stop tools — generic IDEs hit a ceiling on complete games. "
        "But the fourth shift is the one I would like to leave you with. From AI-assisted game making to systematic "
        "expression as a literacy for the AI era. When making a game is easy, the scarce skill is asking a precise question "
        "of an experience. Game design is a natural training ground, and a teachable, accountable one, if the workflow "
        "is structured for it."
    ),
    16: (
        "When making a game becomes easy, the hard part is knowing what you want players to feel. "
        "AI cannot do that for you — but it can help you ask it clearly. Thank you. "
        "I am happy to take questions on the method, the data, the cases, or the next step."
    ),
}


def repair_notes_master(prs: Presentation):
    """The slidep-generated deck ships a notesMaster with an EMPTY spTree, so
    `notes_slide.notes_text_frame` returns None. Transplant the body/slide-image/
    slide-number placeholders from python-pptx's default notesMaster, rescaling
    Y coordinates from the 4:3 notes page (6858000 x 9144000) to this deck's
    16:9 notes page (6858000 x 12192000)."""
    master = prs.notes_master.element
    sp_tree = master.find(qn("p:cSld")).find(qn("p:spTree"))
    if sp_tree.find(qn("p:sp")) is not None:
        return False

    template = Presentation().notes_master.element
    new_tree = copy.deepcopy(template.find(qn("p:cSld")).find(qn("p:spTree")))

    ratio = 12192000 / 9144000
    for tag in ("a:off", "a:ext"):
        for el in new_tree.iter(qn(tag)):
            y = el.get("y")
            if y is not None:
                el.set("y", str(int(round(int(y) * ratio))))

    master.find(qn("p:cSld")).replace(sp_tree, new_tree)
    return True


def add_notes(slide, text: str):
    tf = slide.notes_slide.notes_text_frame
    tf.text = ""
    paragraphs = [p for p in text.split("\n") if p.strip()]
    tf.paragraphs[0].text = paragraphs[0]
    for line in paragraphs[1:]:
        tf.add_paragraph().text = line


def strip_existing_movies(slide) -> int:
    """Make the injection idempotent: drop any previously embedded video so a
    re-run can never stack two copies on slide 12."""
    doomed = [
        sh for sh in slide.shapes
        if sh.shape_type == 16  # MSO_SHAPE_TYPE.MEDIA
        or sh.element.find(f".//{qn('a:videoFile')}") is not None
    ]
    for sh in doomed:
        slide.shapes._spTree.remove(sh.element)
    return len(doomed)


def main():
    prs = Presentation(str(PPTX))
    assert len(prs.slides) == 16, f"Expected 16 slides, found {len(prs.slides)}"
    if repair_notes_master(prs):
        print("[notes master] placeholders restored")
    slide12 = prs.slides[11]
    removed = strip_existing_movies(slide12)
    if removed:
        print(f"[video] removed {removed} pre-existing movie shape(s) on slide 12")
    slide12.shapes.add_movie(
        str(VIDEO),
        LEFT, TOP, WIDTH, HEIGHT,
        poster_frame_image=str(POSTER),
        mime_type="video/mp4",
    )
    print("[video] embedded on slide 12")

    for idx, slide in enumerate(prs.slides, start=1):
        text = NOTES.get(idx, "")
        if text:
            add_notes(slide, text)
    print(f"[notes] added to {len(NOTES)} slides")

    prs.save(str(PPTX))
    print(f"[saved] {PPTX}")


if __name__ == "__main__":
    main()
